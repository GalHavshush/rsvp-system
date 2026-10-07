import csv
import io

MAX_BYTES = 5_000_000
MAX_ROWS = 5000
MAX_COLS = 60


class ParseError(Exception):
    def __init__(self, code: str):
        self.code = code


def _cell(v: object) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))  # spreadsheet phone numbers often arrive as floats
    return str(v).strip()


def _xlsx_rows(data: bytes) -> list[list[str]]:
    from openpyxl import load_workbook

    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)  # formulas are never evaluated
        return [[_cell(c) for c in row[:MAX_COLS]] for row in wb.worksheets[0].iter_rows(values_only=True)]
    except Exception:
        raise ParseError("unreadable_file")


def _csv_rows(data: bytes) -> list[list[str]]:
    for enc in ("utf-8-sig", "cp1255"):  # cp1255 = Hebrew Excel "CSV" export
        try:
            text = data.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise ParseError("unreadable_file")
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    return [[c.strip() for c in row[:MAX_COLS]] for row in csv.reader(io.StringIO(text), dialect)]


def parse_upload(filename: str, data: bytes) -> tuple[list[str], list[list[str]]]:
    """Return (headers, rows). Rows are padded to the header width; blank rows dropped."""
    if len(data) > MAX_BYTES:
        raise ParseError("file_too_large")
    name = filename.lower()
    if name.endswith((".xlsx", ".xlsm")):
        rows = _xlsx_rows(data)
    elif name.endswith(".csv"):
        rows = _csv_rows(data)
    else:
        raise ParseError("unsupported_file_type")
    rows = [r for r in rows if any(r)]
    if not rows:
        raise ParseError("empty_file")
    if len(rows) - 1 > MAX_ROWS:
        raise ParseError("too_many_rows")
    headers, body = rows[0], rows[1:]
    width = len(headers)
    return headers, [(r + [""] * width)[:width] for r in body]
